"""Build the project report from reviewed Vietnamese prose; sources stay intact."""
from __future__ import annotations

import hashlib
import json
import os
import re
import argparse
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PROJECT_TITLE = 'HỆ THỐNG GIÁM SÁT VÀ PHÁT HIỆN SỰ MẤT TẬP TRUNG CỦA TÀI XẾ'
SOURCE = ROOT / "reports/BAO_CAO_GIAM_SAT_BUON_NGU_HOAN_THIEN.md"
TARGET = SOURCE.with_suffix(".docx")
QA = ROOT / "reports/reference_style"
ASSETS = ROOT / "report_assets/project_report"
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".cache/matplotlib"))


def flow_diagram():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

    ASSETS.mkdir(exist_ok=True)
    fig, ax = plt.subplots(figsize=(8.0, 6.2))
    ax.set(xlim=(0, 10), ylim=(0, 10))
    ax.axis("off")
    fig.subplots_adjust(left=0.01, right=0.99, bottom=0.01, top=0.99)

    def box(x, y, w, h, text):
        ax.add_patch(FancyBboxPatch((x-w/2, y-h/2), w, h,
                     boxstyle="round,pad=0.03,rounding_size=0.10",
                     linewidth=1.1, edgecolor="#303030", facecolor="#f6f6f6"))
        ax.text(x, y, text, ha="center", va="center", fontsize=10.8,
                fontfamily="DejaVu Sans", color="#151515")

    def arrow(points, label=None):
        if len(points) > 2:
            ax.plot(*zip(*points[:-1]), color="#303030", linewidth=1.1)
        ax.add_patch(FancyArrowPatch(points[-2], points[-1], arrowstyle="-|>",
                                     mutation_scale=12, color="#303030", linewidth=1.1))
        if label:
            (x, y), s = label
            ax.text(x, y, s, fontsize=10, ha="center", color="#303030")

    box(5, 9.3, 7.2, .9, "Bắt đầu: cấu hình → manifest → kiểm tra model")
    box(5, 7.9, 7.2, .9, "Mở nguồn → đọc frame và timestamp")
    box(5, 6.5, 7.2, .9, "YuNet → chọn một mặt → ROI → chất lượng")
    box(2.7, 4.9, 4.0, 1.0, "Quan sát hợp lệ\nClassifiers → xác suất theo tên lớp")
    box(7.7, 4.9, 3.4, 1.0, "Thiếu hoặc không đạt\nQuan sát UNKNOWN")
    box(5, 3.3, 7.2, .9, "Monitor: trạng thái + lý do + sự kiện theo thời gian")
    box(5, 1.9, 7.2, .9, "Hiển thị (nếu có) + serial + nhật ký")
    box(5, .5, 7.2, .9, "Kết thúc: gửi tắt → đóng tài nguyên → summary")
    for a,b in [(9.3,7.9),(7.9,6.5),(3.3,1.9),(1.9,.5)]:
        arrow([(5,a-.5),(5,b+.5)])
    arrow([(3.3,6.0),(3.3,5.7),(2.7,5.7),(2.7,5.45)])
    arrow([(7.0,6.0),(7.0,5.7),(7.7,5.7),(7.7,5.45)])
    arrow([(2.7,4.35),(2.7,3.95),(4,3.95),(4,3.8)])
    arrow([(7.7,4.35),(7.7,3.95),(6,3.95),(6,3.8)])
    arrow([(1.35,1.9),(.3,1.9),(.3,7.9),(1.3,7.9)], ((.8,8.45),"Frame tiếp"))
    fig.savefig(ASSETS / "runtime_flow.png", dpi=220, facecolor="white")
    plt.close(fig)


def field(p, instruction, cached=""):
    for tag, text in [("begin",None),(None,instruction),("separate",None),(None,cached),("end",None)]:
        r = OxmlElement("w:r")
        if tag:
            child = OxmlElement("w:fldChar")
            child.set(qn("w:fldCharType"), tag)
            if tag == "begin":
                child.set(qn("w:dirty"), "true")
        else:
            child = OxmlElement("w:instrText" if text == instruction else "w:t")
            child.set(qn("xml:space"), "preserve")
            child.text = text
        r.append(child)
        p._p.append(r)


def font(s, size, bold=False, italic=False):
    s.font.name = "Times New Roman"
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = italic
    s.font.color.rgb = RGBColor(0, 0, 0)
    rf = s.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(rf.attrib):
        if "theme" in key.lower():
            del rf.attrib[key]
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:"+name), "Times New Roman")


def configure_styles(doc):
    normal = doc.styles["Normal"]
    font(normal,13)
    p = normal.paragraph_format
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.first_line_indent = Cm(.75)
    p.line_spacing = 1.3
    p.space_after = Pt(6)
    p.widow_control = True
    for level, size in [(1,16),(2,13),(3,13)]:
        s = doc.styles[f"Heading {level}"]
        font(s, size, True)
        p = s.paragraph_format
        p.first_line_indent = Cm(0)
        p.left_indent = Cm(0)
        p.space_before = Pt(12 if level>1 else 0)
        p.space_after = Pt(6 if level>1 else 14)
        p.line_spacing = 1.2
        p.keep_with_next = True
        p.keep_together = True
        p.page_break_before = level == 1
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    for name,size,bold,italic in [
        ("Front Heading",16,True,False), ("Front Subheading",13,True,False),
        ("Report Table",11.5,False,False), ("Report Caption",12,False,True),
        ("Report Code",10.0,False,False),("Report Reference",11.5,False,False),
        ("Report Equation",12,False,False)]:
        s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = normal
        font(s,size,bold,italic)
        p=s.paragraph_format
        p.first_line_indent = Cm(0)
        p.line_spacing = 1.1 if name in ("Report Table","Report Code","Report Reference") else 1.2
        p.space_before = Pt(0)
        p.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if name in ("Front Heading","Report Caption","Report Equation"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if name == "Front Heading":
            p.page_break_before = True
            p.space_after = Pt(14)
        if name in ("Front Subheading","Front Heading"):
            p.keep_with_next = True
            p.space_before = Pt(10)
        if name == "Report Reference":
            p.space_after = Pt(8)
            p.left_indent = Cm(.6)
            p.first_line_indent = Cm(-.6)
    code = doc.styles["Report Code"]
    code.font.name = "Consolas"
    code.element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code.element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    font(doc.styles["Title"],18,True)
    tp=doc.styles["Title"].paragraph_format
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.first_line_indent = Cm(0)
    tp.line_spacing = 1.15
    tp.space_after = Pt(18)
    for border in list(doc.styles["Title"].element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    for level in (1,2,3):
        name = f"TOC {level}"
        s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        s.base_style=normal
        font(s,12)
        p=s.paragraph_format
        p.first_line_indent=Cm(0)
        p.left_indent=Cm((level-1)*.45)
        p.line_spacing=1.05
        p.space_before=Pt(0)
        p.space_after=Pt(4)
        p.keep_with_next=False
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT


def set_page(section, number_format, start, cover=False):
    section.page_width, section.page_height = Cm(21),Cm(29.7)
    section.top_margin,section.bottom_margin = Cm(2.5),Cm(2.5)
    section.left_margin,section.right_margin = Cm(2.5),Cm(2.5)
    section.header_distance,section.footer_distance = Cm(1.0),Cm(1.25)
    num=OxmlElement("w:pgNumType")
    num.set(qn("w:fmt"),number_format)
    num.set(qn("w:start"),str(start))
    section._sectPr.append(num)
    section.footer.is_linked_to_previous = False
    p=section.footer.paragraphs[0]
    p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent=Cm(0)
    if not cover:
        field(p," PAGE ","1")


def cover(doc):
    p=doc.add_paragraph("", "Normal")
    p.paragraph_format.space_after=Pt(2)
    for text in ("TRƯỜNG:","KHOA:","NGÀNH:"):
        p=doc.add_paragraph(text)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.space_after=Pt(9)
    p=doc.add_paragraph("BÁO CÁO ĐỒ ÁN", "Front Heading")
    p.paragraph_format.page_break_before=False
    p.paragraph_format.space_before=Pt(45)
    p.paragraph_format.space_after=Pt(24)
    p=doc.add_paragraph(PROJECT_TITLE, "Title")
    p=doc.add_paragraph("Hướng đến triển khai trên Radxa ROCK 4D")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.space_after=Pt(34)
    for text in ("Sinh viên thực hiện:","Mã số sinh viên:","Lớp:","Giảng viên hướng dẫn:"):
        p=doc.add_paragraph(text)
        p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.left_indent=Cm(2)
        p.paragraph_format.space_after=Pt(13)
    p=doc.add_paragraph("Tháng             năm")
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.space_before=Pt(38)


def table(doc, rows, caption):
    if caption:
        p=doc.add_paragraph(caption,"Report Caption")
        p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(rows[0]))
    t.style="Table Grid"
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    n=len(rows[0])
    if n == 2:
        widths=[5.2,10.8]
    elif n==3:
        widths=[6.6,3.3,6.1] if ("Số mẫu" in rows[0]) else [5.7,5.1,5.2]
    else:
        widths=[6.0]+[10.0/(n-1)]*(n-1)
    for c,w in zip(t.columns,widths):
        c.width=Cm(w)
    for row_index,values in enumerate(rows):
        cells=t.rows[0].cells if row_index==0 else t.add_row().cells
        tr=t.rows[row_index]._tr
        trpr=tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if row_index==0:
            trpr.append(OxmlElement("w:tblHeader"))
        for cell,value,w in zip(cells,values,widths):
            cell.width=Cm(w)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]
            p.style=doc.styles["Report Table"]
            p.paragraph_format.keep_with_next = row_index < len(rows)-1
            r=p.add_run(value)
            if row_index==0:
                r.bold=True
                shade=OxmlElement("w:shd")
                shade.set(qn("w:fill"),"EDEDED")
                cell._tc.get_or_add_tcPr().append(shade)
    p=doc.add_paragraph("")
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=1.0
    p.paragraph_format.first_line_indent=Cm(0)
    return t


def image_paragraph(doc,caption,path):
    from PIL import Image
    image_path=(SOURCE.parent/path).resolve()
    with Image.open(image_path) as im:
        w,h=im.size
    width=min(16.0,15.8*w/h)
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_after=Pt(4)
    p.add_run().add_picture(str(image_path),width=Cm(width))
    captionp=doc.add_paragraph(caption,"Report Caption")
    captionp.paragraph_format.space_after=Pt(10)


def blank_headings(lines):
    entries=[]
    for i,line in enumerate(lines):
        m=re.match(r"^(#{1,3}) (.*)$",line)
        if not m or line.endswith("MỤC LỤC") or line.endswith("BIỂU") or line.endswith("ẢNH"):
            continue
        j=i+1
        contents=[]
        while j<len(lines) and not re.match(r"^#{1,3} ",lines[j]):
            if lines[j].strip(): contents.append(lines[j].strip())
            j+=1
        if not contents:
            entries.append(m.group(2))
    return entries


def build():
    QA.mkdir(parents=True, exist_ok=True)
    flow_diagram()
    lines=SOURCE.read_text(encoding="utf-8").splitlines()
    blank=blank_headings(lines)
    d=Document()
    configure_styles(d)
    set_page(d.sections[0],"decimal",1,True)
    cover(d)
    from docx.enum.section import WD_SECTION_START
    s=d.add_section(WD_SECTION_START.NEW_PAGE)
    set_page(s,"lowerRoman",1)
    d.core_properties.title=PROJECT_TITLE
    d.core_properties.subject="Báo cáo đồ án theo văn phong tài liệu tham khảo"
    d.core_properties.author=""
    d.core_properties.keywords="YuNet, YOLOv8n classification, Radxa ROCK 4D, ESP32"
    d.settings.element.append(OxmlElement("w:updateFields"))
    d.settings.element[-1].set(qn("w:val"),"true")
    started_body=False
    first_heading=True
    i=0
    table_caption=""
    tables=[]
    figures=[]
    while i<len(lines):
        line=lines[i].strip()
        i+=1
        if not line: continue
        m=re.match(r"^(#{1,3}) (.*)$",line)
        if m:
            level=len(m.group(1)); title=m.group(2)
            if title.startswith("CHƯƠNG 1."):
                started_body=True
                s=d.add_section(WD_SECTION_START.NEW_PAGE)
                set_page(s,"decimal",1)
            style=f"Heading {level}" if started_body else "Front Heading" if level==1 else "Front Subheading"
            p=d.add_paragraph(title,style)
            if first_heading or title.startswith("CHƯƠNG 1."):
                p.paragraph_format.page_break_before=False
                first_heading=False
            if title in blank:
                following=next((re.match(r"^(#{1,3}) ",x) for x in lines[i:] if x.strip()),None)
                has_children=following is not None and len(following.group(1))>level
                p.paragraph_format.keep_with_next=has_children
                if not has_children:
                    empty=d.add_paragraph("")
                    empty.paragraph_format.line_spacing=1.0
                    empty.paragraph_format.space_after=Pt(0)
            continue
        if line=="{{TOC}}":
            p=d.add_paragraph("")
            p.paragraph_format.first_line_indent=Cm(0)
            field(p,' TOC \\o "1-3" \\h \\z \\u ',"")
            continue
        if line in ("{{TABLES}}","{{FIGURES}}"):
            label="Bảng" if line=="{{TABLES}}" else "Hình"
            p=d.add_paragraph("")
            p.paragraph_format.first_line_indent=Cm(0)
            field(p,f' TOC \\h \\z \\c "{label}" ',"")
            # Actual figures have literal numbering; populate lists in a later pass.
            p._p.set(qn("w:rsidR"),"00ABC001" if label=="Bảng" else "00ABC002")
            continue
        if line.startswith("@table "):
            table_caption=line[len("@table "):]
            tables.append(table_caption)
            continue
        if line.startswith("|"):
            rows=[]
            while True:
                row=[x.strip() for x in line.strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",x) for x in row): rows.append(row)
                if i>=len(lines) or not lines[i].strip().startswith("|"): break
                line=lines[i].strip();i+=1
            table(d,rows,table_caption)
            table_caption=""
            continue
        if line=="@code":
            codes=[]
            while i<len(lines) and lines[i].strip()!="@endcode":
                codes.append(lines[i]);i+=1
            i+=1
            p=d.add_paragraph("\n".join(codes),"Report Code")
            p.paragraph_format.keep_together=True
            p.paragraph_format.keep_with_next=True
            continue
        if line.startswith("@equation "):
            d.add_paragraph(line[len("@equation "):],"Report Equation")
            continue
        m=re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)",line)
        if m:
            figures.append(m.group(1))
            image_paragraph(d,m.group(1),m.group(2))
            continue
        # Avoid unbreakable long hashes/URLs, while retaining their literal values.
        p=d.add_paragraph(line,"Report Reference" if re.match(r"^\[\d+\]",line) else "Normal")
        if max(map(len,line.split()),default=0)>55:
            p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if line.startswith("Ví dụ "):
            p.style=d.styles["Report Caption"]

    # A list based on captions and PAGEREF bookmarks remains stable after Word updates.
    captions=tables+figures
    bookmark_id=10
    for p in list(d.paragraphs):
        if p.text in captions:
            bookmark_id+=1
            name="cap_"+str(bookmark_id)
            start=OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"),str(bookmark_id));start.set(qn("w:name"),name)
            end=OxmlElement("w:bookmarkEnd");end.set(qn("w:id"),str(bookmark_id))
            p._p.insert(0,start);p._p.append(end)
            for entry in list(d.paragraphs):
                marker=entry._p.get(qn("w:rsidR"))
                if marker == ("00ABC001" if p.text.startswith("Bảng ") else "00ABC002"):
                    # Insert before the marker, in caption occurrence order.
                    row=OxmlElement("w:p")
                    entry._p.addprevious(row)
                    from docx.text.paragraph import Paragraph
                    q=Paragraph(row,entry._parent)
                    q.style=d.styles["Report Reference"]
                    q.paragraph_format.left_indent=Cm(0)
                    q.paragraph_format.first_line_indent=Cm(0)
                    tabs=q.paragraph_format.tab_stops
                    from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
                    tabs.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
                    q.add_run(p.text+"\t")
                    field(q,f' PAGEREF {name} \\h ',"")
                    break
    for p in list(d.paragraphs):
        if p._p.get(qn("w:rsidR")) in ("00ABC001","00ABC002"):
            p._p.getparent().remove(p._p)

    d.save(TARGET)
    # Direct XML/structure check; separate from rendering.
    check=Document(TARGET)
    checks=[]
    for title in blank:
        matches=[p for p in check.paragraphs if p.text==title]
        if len(matches)!=1: raise ValueError("Missing/duplicate blank heading: "+title)
        node=matches[0]._p.getnext()
        text=[]
        while node is not None:
            props=node.find(qn("w:pPr"))
            style=props.find(qn("w:pStyle")) if props is not None else None
            if style is not None and style.get(qn("w:val")) in ("Heading1","Heading2","Heading3","FrontHeading","FrontSubheading"):
                break
            text.extend(x.text or "" for x in node.iter(qn("w:t")))
            node=node.getnext()
        if any(x.strip() for x in text): raise ValueError("Blank section contains content: "+title)
        checks.append(title)
    references={}
    for path in [
        SOURCE,ROOT/"configs/camera_pc.yaml",ROOT/"models/yawn_camera_refined/manifest.json",
        ROOT/"models/eye_refined_v4/manifest.json",ROOT/"models/yawn_camera_refined/evaluation.json",
        ROOT/"reports/yawn_camera_check/candidate.regression.json",
        ROOT/"reports/eye_classifier_verification.json",ROOT/"models/smile_refined/evaluation.json",
        Path(r"C:\Users\P1 Gen 5\Downloads\ung-dung-xu-ly-anh-iot-trong-giam-sat-tai-xe-tren-o-to.pdf"),
        ROOT/"BAO_CAO_DA2_HE_THONG_GIAM_SAT_TAI_XE.docx",
    ]:
        references[str(path)]=hashlib.sha256(path.read_bytes()).hexdigest()
    evidence={"output":str(TARGET),"source_hashes":references,"blank_headings":checks,
              "chapter_count":6,"tables":tables,"figures":figures,
              "paragraph_count":len(check.paragraphs),"word_count":len(SOURCE.read_text(encoding="utf-8").split()),
              "rendered":False}
    (QA/"build_checks.json").write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({k:v for k,v in evidence.items() if k not in ("source_hashes","blank_headings","tables","figures")},ensure_ascii=False))
    print("Blank headings:",len(checks),"Tables:",len(tables),"Figures:",len(figures))


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source",type=Path,default=SOURCE)
    parser.add_argument("--output",type=Path,default=TARGET)
    parser.add_argument("--qa-dir",type=Path,default=QA)
    args=parser.parse_args()
    SOURCE=args.source.resolve()
    TARGET=args.output.resolve()
    QA=args.qa_dir.resolve()
    build()
