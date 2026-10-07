"""Create report diagrams and a formatted DOCX from the reviewed Markdown."""
from __future__ import annotations

import json
import os
import re
from pathlib import Path
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "report_assets/model_architecture"
SOURCE = ROOT / "reports/BAO_CAO_KIEN_TRUC_MODEL.md"
TARGET = SOURCE.with_suffix(".docx")
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".cache/matplotlib"))


def diagrams():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

    ASSETS.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "svg.fonttype": "none"})
    navy, line = "#15324f", "#58718a"
    palette = {"cpu": "#e9f2fb", "model": "#e6f4f0", "logic": "#fff0d9", "plain": "#f2f5f8"}

    def canvas(width, height, limits):
        fig, ax = plt.subplots(figsize=(width, height))
        fig.patch.set_facecolor("white")
        ax.set(xlim=(limits[0], limits[1]), ylim=(limits[2], limits[3]))
        ax.axis("off")
        fig.subplots_adjust(left=0.015, right=0.985, bottom=0.025, top=0.98)
        return fig, ax

    def box(ax, x, y, width, height, label, kind="cpu", fontsize=12):
        ax.add_patch(FancyBboxPatch((x - width / 2, y - height / 2), width, height,
                                   boxstyle="round,pad=0.05,rounding_size=0.09",
                                   linewidth=1.2, edgecolor=line, facecolor=palette[kind], zorder=3))
        ax.text(x, y, label, color=navy, fontsize=fontsize, ha="center", va="center", zorder=4)

    def arrow(ax, start, end, dashed=False):
        ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=14,
                                    color=line, linewidth=1.5, linestyle="--" if dashed else "-", zorder=2))

    def route(ax, points, dashed=False):
        xs, ys = zip(*points[:-1])
        if len(xs) > 1:
            ax.plot(xs, ys, color=line, linewidth=1.4, linestyle="--" if dashed else "-", zorder=2)
        arrow(ax, points[-2], points[-1], dashed)

    def save(fig, name):
        for extension in ("png", "svg"):
            fig.savefig(ASSETS / f"{name}.{extension}", dpi=220, facecolor="white")
        plt.close(fig)

    fig, ax = canvas(8.5, 6.3, (-0.1, 10.1, -0.3, 10.35))
    box(ax, 5, 9.6, 7.3, 0.95, "Camera / video 640×480\nTimestamp: monotonic / PTS hoặc sidecar")
    box(ax, 5, 8.15, 7.3, 0.95, "YuNet trên CPU · BGR 320×320\nBbox mặt + 5 điểm mốc + score")
    box(ax, 5, 6.7, 7.3, 0.95, "Chọn tài xế trong ROI · Kiểm tra quality\nĐổi tọa độ về ảnh gốc rồi crop")
    arrow(ax, (5, 9.05), (5, 8.72))
    arrow(ax, (5, 7.6), (5, 7.26))
    for x, label in ((1.75, "Crop mắt 1\nRGB 224×224"), (5, "Crop mắt 2\nRGB 224×224"), (8.25, "Crop mặt\nRGB 224×224")):
        box(ax, x, 5.1, 2.7, 0.85, label, "plain", 11.5)
        route(ax, [(5, 6.15), (5, 5.92), (x, 5.92), (x, 5.58)])
    box(ax, 3.375, 3.55, 5.95, 1.1, "YOLOv8n-cls mắt · chung trọng số hai mắt\nHai lượt predict → P(Closed), P(Open)", "model", 11.3)
    box(ax, 8.25, 3.55, 2.7, 1.1, "YOLOv8n-cls ngáp\nP(no_yawn), P(yawn)", "model", 10.5)
    arrow(ax, (1.75, 4.62), (1.75, 4.16))
    arrow(ax, (5, 4.62), (5, 4.16))
    arrow(ax, (8.25, 4.62), (8.25, 4.16))
    box(ax, 5, 1.95, 9.1, 1.0, "Monitor: xác suất + quality + thời gian\nNORMAL / WARNING / ALERT / UNKNOWN", "logic", 12)
    route(ax, [(3.375, 2.95), (3.375, 2.67), (5, 2.67), (5, 2.5)])
    route(ax, [(8.25, 2.95), (8.25, 2.67), (5, 2.67), (5, 2.5)])
    box(ax, 2.65, 0.35, 4.4, 0.8, "ESP32 / còi\nSerial + ACK + watchdog", "plain", 11)
    box(ax, 7.4, 0.35, 4.3, 0.8, "Nhật ký JSONL\nTrạng thái · quality · FPS · latency", "plain", 10.5)
    route(ax, [(5, 1.39), (5, 1.08), (2.65, 1.08), (2.65, 0.8)])
    route(ax, [(5, 1.39), (5, 1.08), (7.4, 1.08), (7.4, 0.8)])
    save(fig, "system_pipeline")

    fig, ax = canvas(8.5, 5.8, (-0.1, 10.1, -0.4, 8.45))
    stages = [(7.75, "Input RGB · 3×224×224"),
              (6.65, "Conv 3×3, s=2\n16×112×112"),
              (5.55, "Conv 3×3, s=2 → C2f(n=1)\n32×56×56"),
              (4.45, "Conv 3×3, s=2 → C2f(n=2)\n64×28×28"),
              (3.35, "Conv 3×3, s=2 → C2f(n=2)\n128×14×14"),
              (2.25, "Conv 3×3, s=2 → C2f(n=1)\n256×7×7")]
    for index, (y, label) in enumerate(stages):
        box(ax, 3, y, 5.15, 0.8, label, "plain" if index == 0 else "cpu", 12)
        if index:
            arrow(ax, (3, stages[index - 1][0] - 0.46), (3, y + 0.46))
    box(ax, 8.08, 5.8, 3.0, 3.05,
        "Mỗi classifier\n1.440.850 tham số\n0,2013 GMAC/lượt\n\nMắt/ngáp cùng cấu trúc,\ntrọng số riêng.\n\nTensor: C×H×W", "plain", 11.4)
    box(ax, 8.08, 2.25, 3.0, 0.8, "Conv 1×1: 256→1280\n1280×7×7", "model", 10.7)
    arrow(ax, (5.62, 2.25), (6.49, 2.25))
    box(ax, 8.08, 1.15, 3.0, 0.8, "GAP → Flatten\nVector 1280 chiều", "model", 11)
    arrow(ax, (8.08, 1.79), (8.08, 1.61))
    box(ax, 8.08, 0.05, 3.0, 0.8, "Linear 1280→2 + Softmax\n[1,2] xác suất", "logic", 10.5)
    arrow(ax, (8.08, 0.69), (8.08, 0.51))
    ax.text(2.95, 0.7, "Backbone: Conv + C2f\nHead: Conv → GAP → Linear → Softmax", fontsize=11.5,
            color=navy, ha="center", va="center")
    save(fig, "classifier_architecture")

    fig, ax = canvas(8.5, 4.0, (-0.1, 12.4, -0.1, 5.35))
    box(ax, 0.8, 2.8, 1.15, 0.95, "X\n64×28×28", "plain", 10)
    box(ax, 2.45, 2.8, 1.45, 0.95, "Conv 1×1\n64→64", "cpu", 11)
    arrow(ax, (1.45, 2.8), (1.64, 2.8))
    box(ax, 4.25, 4.5, 1.25, 0.8, "A\n32 kênh", "cpu", 11)
    box(ax, 4.25, 2.8, 1.25, 0.8, "B\n32 kênh", "cpu", 11)
    route(ax, [(3.25, 2.8), (3.43, 2.8), (3.43, 4.5), (3.55, 4.5)])
    arrow(ax, (3.25, 2.8), (3.55, 2.8))
    box(ax, 6.05, 2.8, 1.4, 0.95, "Bottleneck 1\n32 kênh", "logic", 10.3)
    box(ax, 7.95, 2.8, 1.4, 0.95, "Bottleneck 2\n32 kênh", "logic", 10.3)
    arrow(ax, (4.94, 2.8), (5.3, 2.8))
    arrow(ax, (6.81, 2.8), (7.2, 2.8))
    box(ax, 9.7, 2.8, 1.18, 2.05, "Concat\nA, B, B1, B2\n128 kênh", "model", 9.5)
    route(ax, [(4.94, 4.5), (8.85, 4.5), (8.85, 3.52), (9.04, 3.52)])
    route(ax, [(4.96, 2.8), (5.08, 2.8), (5.08, 0.7), (8.87, 0.7), (8.87, 2.05), (9.04, 2.05)])
    route(ax, [(6.84, 2.8), (7.01, 2.8), (7.01, 1.35), (8.66, 1.35), (8.66, 2.38), (9.04, 2.38)])
    arrow(ax, (8.73, 2.8), (9.04, 2.8))
    box(ax, 11.55, 2.8, 1.35, 0.95, "Conv 1×1\n128→64", "cpu", 11)
    arrow(ax, (10.34, 2.8), (10.8, 2.8))
    ax.text(6.2, 5.1, "C2f tầng 4: 64 kênh vào/ra · n=2 · H=W=28", color=navy, fontsize=13, ha="center")
    ax.text(6.2, 0.08, "Mỗi Bottleneck: hai Conv 3×3 + shortcut · Concatenate theo chiều kênh", color=navy,
            fontsize=10.2, ha="center")
    save(fig, "c2f_block")


def create_docx():
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    inventory = json.loads((ROOT / "reports/model_architecture_inventory.json").read_text(encoding="utf-8"))
    manifest = json.loads((ROOT / "models/manifest.json").read_text(encoding="utf-8"))
    for name, item in inventory["classifiers"].items():
        if item["torch"]["total_parameters"] != 1_440_850 or item["torch"]["conv_linear_macs_at_batch1"] != 201_308_672:
            raise ValueError("Architecture changed: update the reviewed report before generating DOCX")
        for kind in ("torch", "onnx", "rknn"):
            if item[kind]["sha256"] != manifest["classifiers"][name]["artifacts"][kind]["sha256"]:
                raise ValueError(f"Inventory is stale for {name}/{kind}")
    markdown = SOURCE.read_text(encoding="utf-8")
    for spec in manifest["classifiers"].values():
        for item in spec["artifacts"].values():
            if item["sha256"] not in markdown:
                raise ValueError("A model hash is missing in the reviewed report")
    if manifest["detector"]["sha256"] not in markdown:
        raise ValueError("YuNet hash missing")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2), Cm(2)
    section.header_distance, section.footer_distance = Cm(0.7), Cm(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 21), ("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string("15324F")
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    doc.styles["Caption"].font.name = "Times New Roman"
    doc.styles["Caption"].font.size = Pt(10)
    doc.styles["Caption"].font.italic = True
    doc.styles["Caption"].font.color.rgb = RGBColor.from_string("56687A")
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("GIÁM SÁT TÀI XẾ  |  KIẾN TRÚC MODEL")
    run.font.name, run.font.size = "Calibri", Pt(8.5)
    run.font.color.rgb = RGBColor.from_string("56687A")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("06/10/2026   •   Trang ").font.size = Pt(9)
    for field, ending in (("PAGE", " / "), ("NUMPAGES", "")):
        field_node = OxmlElement("w:fldSimple")
        field_node.set(qn("w:instr"), field)
        footer._p.append(field_node)
        footer.add_run(ending).font.size = Pt(9)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    doc.core_properties.title = "Báo cáo kiến trúc model giám sát buồn ngủ tài xế"
    doc.core_properties.subject = "YuNet, hai YOLOv8n-cls và runtime Radxa ROCK 4D"
    doc.core_properties.author = "Driver Monitor Project"
    doc.core_properties.keywords = "YuNet;YOLOv8n-cls;RK3576;RKNN;giám sát tài xế"
    doc.core_properties.comments = "Phân tích từ checkpoint và mã nguồn hiện có ngày 06/10/2026."

    token_pattern = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")

    def inline(paragraph, content, size=None):
        for token in token_pattern.split(content):
            if not token:
                continue
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if match:
                label, target = match.groups()
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
                node = OxmlElement("w:r")
                props = OxmlElement("w:rPr")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "176789")
                props.append(color)
                underline = OxmlElement("w:u")
                underline.set(qn("w:val"), "single")
                props.append(underline)
                if size:
                    point = OxmlElement("w:sz")
                    point.set(qn("w:val"), str(int(size * 2)))
                    props.append(point)
                node.append(props)
                text_node = OxmlElement("w:t")
                text_node.text = label
                node.append(text_node)
                link.append(node)
                paragraph._p.append(link)
            else:
                bold = token.startswith("**") and token.endswith("**")
                code = token.startswith("`") and token.endswith("`")
                text = token[2:-2] if bold else token[1:-1] if code else token
                run = paragraph.add_run(text)
                run.bold = bold
                if size:
                    run.font.size = Pt(size)
                if code:
                    run.font.name = "Consolas"
                    run.font.size = Pt(min(size or 10, 10))

    widths = {
        "Thành phần": [2.3, 4.6, 4.5, 5.6],
        "Stride": [1.3, 3.5, 2.0, 3.4, 3.4, 3.4],
        "Vùng ảnh": [3.2, 6.3, 7.5],
        "Tiêu chí mặc định": [6.4, 4.0, 6.6],
        "Tầng": [1.0, 2.05, 5.0, 5.15, 3.8],
        "Checkpoint": [3.5, 3.3, 3.3, 6.9],
        "Tham số": [6.3, 5.35, 5.35],
        "Dấu hiệu / tình huống": [4.7, 7.5, 4.8],
        "Công việc": [5.2, 5.9, 5.9],
        "File model": [6.6, 4.6, 5.8],
        "Phép đo đã lưu": [4.6, 5.6, 6.8],
        "Artifact": [4.0, 13.0],
    }

    def table(rows):
        element = doc.add_table(rows=1, cols=len(rows[0]))
        element.alignment = WD_TABLE_ALIGNMENT.CENTER
        element.autofit = False
        column_widths = widths.get(rows[0][0], [17 / len(rows[0])] * len(rows[0]))
        for column, width in zip(element.columns, column_widths):
            column.width = Cm(width)
        for index, row in enumerate(rows):
            cells = element.rows[0].cells if index == 0 else element.add_row().cells
            for column_index, (cell, value) in enumerate(zip(cells, row)):
                cell.width = Cm(column_widths[column_index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.line_spacing = 1
                inline(paragraph, value, 8.5 if rows[0][0] == "Artifact" and column_index == 1 else 10)
                properties = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "15324F" if index == 0 else "F0F4F8" if index % 2 else "FFFFFF")
                properties.append(shading)
                margins = OxmlElement("w:tcMar")
                vertical_margin = 40 if rows[0][0] == "Tầng" else 65
                for side, amount in (("top", vertical_margin), ("bottom", vertical_margin), ("left", 85), ("right", 85)):
                    margin = OxmlElement(f"w:{side}")
                    margin.set(qn("w:w"), str(amount))
                    margin.set(qn("w:type"), "dxa")
                    margins.append(margin)
                properties.append(margins)
                if index == 0:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            row_properties = element.rows[index]._tr.get_or_add_trPr()
            no_split = OxmlElement("w:cantSplit")
            row_properties.append(no_split)
            if index == 0:
                repeat_header = OxmlElement("w:tblHeader")
                repeat_header.set(qn("w:val"), "true")
                row_properties.append(repeat_header)
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(1)
        spacer.paragraph_format.space_before = Pt(0)
        spacer.add_run().font.size = Pt(2)

    lines, position, pending_break = markdown.splitlines(), 0, False
    while position < len(lines):
        line = lines[position].strip()
        position += 1
        if not line:
            continue
        if line == "<!-- page-break -->":
            pending_break = True
            continue
        if line.startswith("```"):
            while position < len(lines) and not lines[position].startswith("```"):
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(lines[position])
                run.font.name, run.font.size = "Consolas", Pt(9.5)
                position += 1
            position += 1
            continue
        if line.startswith("|"):
            rows = [[piece.strip() for piece in line.strip("|").split("|")]]
            while position < len(lines) and lines[position].strip().startswith("|"):
                cells = [piece.strip() for piece in lines[position].strip().strip("|").split("|")]
                position += 1
                if all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    continue
                if len(cells) != len(rows[0]):
                    raise ValueError("Malformed Markdown table")
                rows.append(cells)
            table(rows)
            continue
        image_match = re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)", line)
        if image_match:
            caption, target = image_match.groups()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            image_width = 15.5 if Path(target).name == "classifier_architecture.png" else 17
            shape = paragraph.add_run().add_picture(str((SOURCE.parent / target).resolve()), width=Cm(image_width))
            shape._inline.docPr.set("descr", caption)
            paragraph = doc.add_paragraph(caption, style="Caption")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        heading = re.match(r"^(#{1,3}) (.*)$", line)
        if heading:
            level, text = len(heading.group(1)), heading.group(2)
            paragraph = doc.add_paragraph(style="Title" if level == 1 else "Heading 1" if level == 2 else "Heading 2")
            inline(paragraph, text)
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            inline(paragraph, line[2:])
        else:
            content = [line]
            while position < len(lines):
                candidate = lines[position].strip()
                if not candidate or candidate.startswith(("#", "|", "![", "- ", "```", "<!--")):
                    break
                content.append(candidate)
                position += 1
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            inline(paragraph, " ".join(content))
        if pending_break:
            paragraph.paragraph_format.page_break_before = True
            pending_break = False
    doc.save(TARGET)
    layer_table = next(t for t in doc.tables if t.cell(0, 0).text == "Tầng")
    for stage in inventory["classifiers"]["eye"]["torch"]["stages"]:
        row = layer_table.rows[stage["index"] + 2]
        recorded_parameters = int(row.cells[4].text.replace(".", ""))
        recorded_shape = row.cells[3].text.replace(" ", "")
        expected_shape = json.dumps(stage["output_shape"], separators=(",", ":"))
        if recorded_parameters != stage["parameters"] or recorded_shape != expected_shape:
            raise ValueError(f"Report layer differs from checkpoint: {stage['index']}")
    if int(layer_table.rows[-1].cells[4].text.replace(".", "")) != 1_440_850:
        raise ValueError("Incorrect total parameter count in report")
    with ZipFile(TARGET) as package:
        if package.testzip() is not None:
            raise ValueError("DOCX ZIP is invalid")
        from lxml import etree
        xml = etree.fromstring(package.read("word/document.xml"))
        namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        extracted = " ".join(xml.xpath("//w:t/text()", namespaces=namespaces))
        if "1.440.850" not in extracted or "RK3576" not in extracted:
            raise ValueError("Required model facts missing from DOCX")
        pictures = [name for name in package.namelist() if name.startswith("word/media/")]
        if len(pictures) != 3:
            raise ValueError("Expected three architecture diagrams")
    print(f"Created {TARGET} ({TARGET.stat().st_size} bytes)")
    print(f"Checked DOCX XML, {len(doc.tables)} tables, {len(pictures)} embedded figures, layer shapes/counts and artifact hashes.")


if __name__ == "__main__":
    diagrams()
    create_docx()
