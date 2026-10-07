"""Verify report scope, source preservation and rendered layout; build QA images."""
import argparse
import hashlib
import io
import json
import re
import sys
import unittest
import zipfile
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
QA = ROOT / "reports/reference_style/system_design_scope"
SOURCE = ROOT / "reports/BAO_CAO_GIAM_SAT_BUON_NGU_THEO_MAU.md"
DOCX = ROOT / "reports/BAO_CAO_GIAM_SAT_BUON_NGU_HOAN_THIEN.docx"
PDF = DOCX.with_suffix(".pdf")
FORBIDDEN = r"\b(?:GPS|SIM800L|GSM|GPRS|SMS|cvzone|FaceNet|AdaBoost|Haar(?:-like)?|EAR|MAR|MQTT|Firebase)\b|giọng nói|danh tính|website|bảng điều khiển từ xa"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(name, value):
    QA.mkdir(parents=True, exist_ok=True)
    (QA / name).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def tests():
    stream = io.StringIO()
    suite = unittest.defaultTestLoader.discover(str(ROOT / "tests"), top_level_dir=str(ROOT))
    result = unittest.TextTestRunner(stream=stream, verbosity=2).run(suite)
    QA.mkdir(parents=True, exist_ok=True)
    (QA / "software_tests.txt").write_text(stream.getvalue(), encoding="utf-8")
    report = {"command": ".venv/Scripts/python.exe -m unittest discover -s tests -v",
              "tests_run": result.testsRun, "failures": len(result.failures), "errors": len(result.errors),
              "passed": result.wasSuccessful(), "physical_hardware_tested": False,
              "active_config_sha256": digest(ROOT / "configs/camera_pc.yaml"),
              "active_manifest_sha256": digest(ROOT / "models/yawn_camera_refined/manifest.json")}
    save("software_tests.json", report)
    print(json.dumps(report), flush=True)
    if not result.wasSuccessful():
        raise RuntimeError("Current software checks failed")


def verify():
    from docx import Document
    from docx.oxml.ns import qn
    from build_project_style_report import blank_headings
    import fitz
    from PIL import Image, ImageDraw

    source = SOURCE.read_text(encoding="utf-8")
    doc = Document(DOCX)
    with zipfile.ZipFile(DOCX) as archive:
        if archive.testzip() is not None:
            raise ValueError("Corrupt DOCX")
    content = "\n".join(p.text for p in doc.paragraphs)
    content += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    if re.search(FORBIDDEN, source + content, re.IGNORECASE):
        raise ValueError("Removed topic persists in report")
    numbered = defaultdict(list)
    for heading in re.findall(r"^#{1,3} (.+)$", source, re.MULTILINE):
        match = re.match(r"(\d+(?:\.\d+)+)\.\s", heading)
        if match:
            parts = tuple(map(int, match.group(1).split('.')))
            numbered[parts[:-1]].append(parts[-1])
    for prefix, numbers in numbered.items():
        if numbers != list(range(1, len(numbers) + 1)):
            raise ValueError(f"Heading sequence invalid: {prefix}: {numbers}")
    blanks = blank_headings(source.splitlines())
    for title in blanks:
        matches = [p for p in doc.paragraphs if p.text == title]
        if len(matches) != 1:
            raise ValueError(f"Missing or duplicate empty heading: {title}")
        node = matches[0]._p.getnext()
        values = []
        while node is not None:
            properties = node.find(qn("w:pPr"))
            style = properties.find(qn("w:pStyle")) if properties is not None else None
            if style is not None and style.get(qn("w:val")) in ("Heading1", "Heading2", "Heading3", "FrontHeading", "FrontSubheading"):
                break
            values.extend(x.text or "" for x in node.iter(qn("w:t")))
            node = node.getnext()
        if any(x.strip() for x in values):
            raise ValueError(f"Unsupported fill in empty heading: {title}")
    figures = re.findall(r"^!\[([^]]+)\]\(([^)]+)\)$", source, re.MULTILINE)
    figure_sources = {caption: digest((SOURCE.parent / path).resolve()) for caption, path in figures}
    # The front matter also includes an unnumbered abbreviations table.
    if len(figures) != 14 or len(doc.inline_shapes) != 14 or len(doc.tables) != 15:
        raise ValueError("Missing figures/tables")
    equations = [p.text for p in doc.paragraphs if p.style.name == "Report Equation"]
    if len(equations) != 20:
        raise ValueError("Algorithm theory equations were lost")
    audit = json.loads((QA / "scope_audit.json").read_text(encoding="utf-8"))
    for path, expected in audit["protected_sources"].items():
        if digest(Path(path)) != expected:
            raise ValueError(f"Reference/source changed: {path}")
    pdf = fitz.open(PDF)
    image_pages, sparse_pages, outside_page, pdf_text, thumbnails = [], [], [], [], []
    previews = QA / "pages"
    previews.mkdir(parents=True, exist_ok=True)
    for index, page in enumerate(pdf):
        text = page.get_text()
        pdf_text.append(text)
        body = '\n'.join(block[4] for block in page.get_text("blocks") if len(block) > 4 and isinstance(block[4], str) and block[1] < page.rect.height - 45)
        if len(body.strip()) < 10:
            sparse_pages.append(index+1)
        for block in page.get_text("dict")["blocks"]:
            x0, y0, x1, y1 = block["bbox"]
            if x0 < -1 or y0 < -1 or x1 > page.rect.width + 1 or y1 > page.rect.height + 1:
                outside_page.append({"page": index+1, "bbox": list(block["bbox"]), "type": block["type"]})
        if page.get_images():
            image_pages.append(index+1)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.0, 1.0), alpha=False)
        thumb = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
        thumb.thumbnail((595, 842))
        thumbnails.append(thumb)
        if page.get_images() or index < 2:
            page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False).save(previews / f"page_{index+1:02d}.png")
    for start in range(0, len(thumbnails), 9):
        sheet = Image.new("RGB", (1815, 2631), "#d7d7d7")
        draw = ImageDraw.Draw(sheet)
        for offset, thumb in enumerate(thumbnails[start:start+9]):
            x, y = (offset % 3) * 605 + 5, (offset // 3) * 877 + 25
            sheet.paste(thumb, (x, y))
            draw.text((x+5, y-20), f"Page {start+offset+1}", fill="black")
        sheet.save(previews / f"contact_{start//9+1:02d}.png")
    combined = '\n'.join(pdf_text)
    if re.search(FORBIDDEN, combined, re.IGNORECASE):
        raise ValueError("Removed topic remains in rendered PDF/TOC")
    if "3.1.3. Lưu đồ hoạt động toàn bộ chương trình" not in combined.replace('\n', ' '):
        raise ValueError("Updated heading not found in PDF")
    if "yawn_camera_refined" not in combined:
        raise ValueError("Active model omitted from PDF")
    if sparse_pages or outside_page:
        raise ValueError(f"Layout failed: sparse={sparse_pages}; outside={outside_page}")
    report = {"output_docx": str(DOCX), "output_pdf": str(PDF), "pages": len(pdf),
              "figures": len(figures), "tables": len(doc.tables), "numbered_tables": 14, "equations": len(equations),
              "empty_headings_preserved": len(blanks), "removed_headings": len(audit["removed_headings"]),
              "out_of_scope_matches": [], "sparse_pages": sparse_pages, "outside_page": outside_page,
              "image_pages": image_pages, "figure_source_sha256": figure_sources,
              "docx_sha256": digest(DOCX), "pdf_sha256": digest(PDF), "protected_sources_unchanged": True,
              "rendered": True, "visual_inspection": "pending; inspect page images before delivery"}
    save("final_checks.json", report)
    print(json.dumps({k:v for k,v in report.items() if k != "figure_source_sha256"}), flush=True)
    pdf.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("stage", choices=("tests", "verify"))
    args = parser.parse_args()
    tests() if args.stage == "tests" else verify()
